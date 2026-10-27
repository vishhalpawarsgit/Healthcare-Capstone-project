
from pathlib import Path
from docx import Document

project_root = Path(__file__).resolve().parents[2]
test_dir = project_root / "data" / "processed" / "ingestion_tests"
test_dir.mkdir(parents=True, exist_ok=True)

# Create a fictional hospital SOP in DOCX format.
docx_path = test_dir / "sample_hospital_sop.docx"
document = Document()
document.add_heading("Fictional Hospital SOP", level=1)
document.add_paragraph(
    "Nursing staff should verify the assigned bed "
    "before a fictional patient admission."
)
document.save(docx_path)

# Create a fictional hospital SOP in TXT format.
txt_path = test_dir / "sample_hospital_sop.txt"
txt_path.write_text(
    "Fictional Hospital SOP\n"
    "Nursing staff should verify the assigned bed "
    "before a fictional patient admission.\n",
    encoding="utf-8",
)

print("Created DOCX:", docx_path)
print("Created TXT:", txt_path)
